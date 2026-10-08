"""PDF -> Word (.docx), dengan OCR untuk halaman hasil scan.

Alur:
1. Setiap halaman dibaca dengan PyMuPDF. Halaman yang punya teks langsung diambil
   (beserta ukuran huruf, tebal, dan miring).
2. Halaman yang hampir tanpa teks tetapi berisi gambar dianggap hasil scan -> OCR (Tesseract).
3. Semua hasil dirangkai menjadi .docx dengan python-docx.

Keterbatasan: layout kolom, tabel, dan gambar tidak ikut diubah (hanya teks).
"""
from __future__ import annotations

import functools
import logging
import os
import re
import shutil
import time
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image

from services.job_context import checkpoint
from services.pdf_common import MissingDependencyError, PdfProcessingError, get_fitz, open_fitz_document
from services.pdf_to_jpg import compute_scale
from services.process_utils import ProcessTimeout, run_command

logger = logging.getLogger(__name__)

MIN_TEXT_CHARS = 20         # halaman dengan teks lebih sedikit dari ini (dan ada gambar) = scan
MAX_OCR_PAGES = 20          # batas halaman OCR per proses (OCR lambat)
OCR_DPI = 200
OCR_MAX_SIDE_PX = 4000
OCR_PAGE_TIMEOUT = 60       # detik per halaman
OCR_TOTAL_TIMEOUT = 300     # detik untuk seluruh OCR
MAX_PAGES = 500

KIND_LABELS = {
    "text": "📄 Text-based PDF (teks diambil langsung)",
    "scanned": "🖼 Scanned PDF (dibaca dengan OCR)",
    "mixed": "📄🖼 Campuran (halaman scan dibaca dengan OCR)",
}

_LANG_RE = re.compile(r"^[A-Za-z0-9_]+(\+[A-Za-z0-9_]+)*$")
_LANG_LINE = re.compile(r"^[A-Za-z0-9_]+$")
# Karakter yang tidak boleh ada di XML (python-docx akan error jika ada).
_BAD_XML_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_TESSERACT_FALLBACKS = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    "/opt/homebrew/bin/tesseract",
    "/usr/local/bin/tesseract",
)


# ---------------------------------------------------------------------------
# Struktur data
# ---------------------------------------------------------------------------
@dataclass
class Span:
    text: str
    size: float = 11.0
    bold: bool = False
    italic: bool = False


@dataclass
class Block:
    """Satu paragraf: rangkaian span dengan format masing-masing."""

    spans: list[Span]

    @property
    def text(self) -> str:
        return "".join(s.text for s in self.spans)


@dataclass
class PageContent:
    width: float = 595.0
    height: float = 842.0
    blocks: list[Block] = field(default_factory=list)
    has_images: bool = False
    note: str | None = None

    @property
    def char_count(self) -> int:
        return sum(len("".join(b.text.split())) for b in self.blocks)

    @property
    def needs_ocr(self) -> bool:
        return self.char_count < MIN_TEXT_CHARS and self.has_images


@dataclass
class ConversionReport:
    kind: str                   # "text" | "scanned" | "mixed"
    page_count: int
    ocr_pages: int = 0
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Ekstraksi teks (PyMuPDF)
# ---------------------------------------------------------------------------
def clean_text(text: str) -> str:
    return _BAD_XML_CHARS.sub("", text)


def _append_span(spans: list[Span], span: Span) -> None:
    last = spans[-1] if spans else None
    if last and (last.size, last.bold, last.italic) == (span.size, span.bold, span.italic):
        last.text += span.text
    else:
        spans.append(span)


def block_from_dict(block: dict) -> Block | None:
    """Ubah satu blok teks hasil `page.get_text("dict")` menjadi Block (satu paragraf)."""
    spans: list[Span] = []
    lines = block.get("lines", [])
    for index, line in enumerate(lines):
        for raw in line.get("spans", []):
            text = clean_text(raw.get("text", ""))
            if not text:
                continue
            flags = int(raw.get("flags", 0))
            bold = bool(flags & 16) or "bold" in str(raw.get("font", "")).lower()
            size = round(float(raw.get("size", 11.0)) * 2) / 2
            _append_span(spans, Span(text, size, bold, bool(flags & 2)))
        # Baris yang dipatahkan di tengah paragraf disambung dengan spasi.
        if index < len(lines) - 1 and spans and not spans[-1].text.endswith(("-", " ")):
            _append_span(spans, Span(" ", spans[-1].size, spans[-1].bold, spans[-1].italic))

    if spans:
        spans[-1].text = spans[-1].text.rstrip()
    spans = [s for s in spans if s.text]
    if not spans or not "".join(s.text for s in spans).strip():
        return None
    return Block(spans)


def extract_pages(input_path: Path) -> list[PageContent]:
    """Baca teks tiap halaman (tanpa OCR)."""
    pages: list[PageContent] = []
    with open_fitz_document(input_path) as doc:
        count = doc.page_count
        if count < 1:
            raise PdfProcessingError("PDF tidak memiliki halaman.")
        if count > MAX_PAGES:
            raise PdfProcessingError(f"PDF terlalu panjang. Maksimal {MAX_PAGES} halaman.")
        for index in range(count):
            checkpoint(index, count, 0.0, 0.2)
            page = doc[index]
            data = page.get_text("dict", sort=True)
            blocks = [
                b for raw in data.get("blocks", [])
                if raw.get("type") == 0 and (b := block_from_dict(raw)) is not None
            ]
            pages.append(
                PageContent(
                    width=float(page.rect.width),
                    height=float(page.rect.height),
                    blocks=blocks,
                    has_images=bool(page.get_images()),
                )
            )
    return pages


def classify(pages: list[PageContent]) -> str:
    """"text" jika tidak ada halaman scan, "scanned" jika semua halaman berisi scan, selain itu "mixed"."""
    scan = sum(1 for p in pages if p.needs_ocr)
    if scan == 0:
        return "text"
    text = sum(1 for p in pages if p.char_count >= MIN_TEXT_CHARS)
    return "scanned" if text == 0 else "mixed"


# ---------------------------------------------------------------------------
# OCR (Tesseract)
# ---------------------------------------------------------------------------
def find_tesseract(configured: str | None = None) -> str | None:
    if configured:
        if Path(configured).is_file():
            return str(configured)
        found = shutil.which(configured)
        if found:
            return found
    found = shutil.which("tesseract")
    if found:
        return found
    for candidate in _TESSERACT_FALLBACKS:
        if Path(candidate).is_file():
            return candidate
    return None


@functools.lru_cache(maxsize=4)
def list_languages(tesseract_cmd: str) -> tuple[str, ...]:
    """Bahasa yang terpasang di Tesseract (hasil `tesseract --list-langs`)."""
    try:
        _code, out, err = run_command([tesseract_cmd, "--list-langs"], timeout=20)
    except (ProcessTimeout, OSError):
        return ()
    text = (out + b"\n" + err).decode("utf-8", "replace")
    langs = [line.strip() for line in text.splitlines() if _LANG_LINE.match(line.strip())]
    return tuple(lang for lang in langs if lang != "osd")


def resolve_languages(requested: str, available: tuple[str, ...]) -> tuple[str, list[str]]:
    """Pilih bahasa OCR yang benar-benar terpasang. Return (kode_bahasa, bahasa_yang_hilang)."""
    wanted = [lang for lang in requested.split("+") if lang]
    usable = [lang for lang in wanted if lang in available]
    missing = [lang for lang in wanted if lang not in available]
    if usable:
        return "+".join(usable), missing
    if "eng" in available:
        return "eng", missing
    if available:
        return available[0], missing
    raise PdfProcessingError("Data bahasa Tesseract tidak ditemukan. Pasang language data (mis. eng).")


def ocr_text_to_paragraphs(text: str) -> list[str]:
    """Pecah keluaran Tesseract menjadi paragraf (dipisah baris kosong)."""
    text = text.replace("\x0c", "").replace("\r\n", "\n")
    paragraphs = [clean_text(" ".join(p.split())) for p in re.split(r"\n\s*\n", text)]
    return [p for p in paragraphs if p]


def run_tesseract(
    image_path: Path, tesseract_cmd: str, languages: str, timeout: float = OCR_PAGE_TIMEOUT
) -> str:
    if not _LANG_RE.match(languages):
        raise ValueError(f"Kode bahasa tidak valid: {languages!r}")
    env = dict(os.environ, OMP_THREAD_LIMIT="1")  # satu thread per proses = lebih stabil
    cmd = [tesseract_cmd, str(image_path), "stdout", "-l", languages, "--psm", "3"]
    try:
        code, out, err = run_command(cmd, timeout=timeout, env=env)
    except ProcessTimeout:
        raise PdfProcessingError("OCR memakan waktu terlalu lama dan dihentikan.") from None
    if code != 0:
        logger.warning("Tesseract gagal (kode %s): %s", code, err[-300:].decode("utf-8", "replace"))
        raise PdfProcessingError("OCR gagal membaca halaman.")
    return out.decode("utf-8", errors="replace")


def ocr_pages(
    input_path: Path,
    indexes: list[int],
    pages: list[PageContent],
    work_dir: Path,
    tesseract_cmd: str,
    languages: str,
) -> None:
    """Jalankan OCR pada halaman `indexes`; hasilnya mengisi `pages[i].blocks`."""
    fitz = get_fitz()
    deadline = time.monotonic() + OCR_TOTAL_TIMEOUT
    image_path = Path(work_dir) / "ocr_page.png"

    with open_fitz_document(input_path) as doc:
        for step, index in enumerate(indexes):
            checkpoint(step, len(indexes), 0.2, 0.9)
            if time.monotonic() > deadline:
                raise PdfProcessingError(
                    "Proses OCR memakan waktu terlalu lama.\n\nCoba PDF dengan halaman lebih sedikit."
                )
            page = doc[index]
            scale = compute_scale(page.rect.width, page.rect.height, OCR_DPI, OCR_MAX_SIDE_PX)
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
            Image.frombytes("RGB", (pix.width, pix.height), pix.samples).convert("L").save(image_path)
            text = run_tesseract(image_path, tesseract_cmd, languages)
            pages[index].blocks = [Block([Span(p)]) for p in ocr_text_to_paragraphs(text)]
    image_path.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# Pembuatan .docx (python-docx)
# ---------------------------------------------------------------------------
def _body_size(pages: list[PageContent]) -> float:
    """Ukuran huruf paling umum (median berbobot panjang teks) = ukuran teks isi."""
    pairs = sorted((s.size, len(s.text)) for p in pages for b in p.blocks for s in b.spans)
    total = sum(w for _, w in pairs)
    if not total:
        return 11.0
    running = 0
    for size, weight in pairs:
        running += weight
        if running >= total / 2:
            return size
    return pairs[-1][0]


def build_docx(pages: list[PageContent], output: Path) -> None:
    doc = Document()
    first = pages[0]
    section = doc.sections[0]
    section.page_width = Pt(min(max(first.width, 200), 2000))
    section.page_height = Pt(min(max(first.height, 200), 2000))
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Inches(0.8))

    body = _body_size(pages)
    for number, page in enumerate(pages):
        if number > 0:
            doc.add_page_break()
        for block in page.blocks:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            is_heading = max(s.size for s in block.spans) >= body * 1.25
            for span in block.spans:
                run = paragraph.add_run(span.text)
                run.bold = span.bold or is_heading
                run.italic = span.italic
                run.font.size = Pt(min(max(span.size, 7), 36))
        if page.note:
            note = doc.add_paragraph().add_run(page.note)
            note.italic = True
    doc.save(str(output))


# ---------------------------------------------------------------------------
# Fungsi utama
# ---------------------------------------------------------------------------
def pdf_to_word(
    input_path: Path,
    output_path: Path,
    work_dir: Path,
    *,
    tesseract_path: str | None = None,
    languages: str = "eng",
) -> ConversionReport:
    """Ubah PDF menjadi DOCX. Halaman hasil scan dibaca dengan OCR bila Tesseract tersedia."""
    pages = extract_pages(input_path)
    kind = classify(pages)
    scan_indexes = [i for i, p in enumerate(pages) if p.needs_ocr]
    warnings: list[str] = []
    ocr_done = 0

    if scan_indexes:
        tesseract = find_tesseract(tesseract_path)
        if tesseract is None:
            logger.error("Tesseract tidak ditemukan. Pasang Tesseract atau isi TESSERACT_PATH di .env")
            if kind == "scanned":
                raise MissingDependencyError(
                    "PDF ini berupa hasil scan dan membutuhkan OCR, "
                    "tetapi Tesseract belum terpasang di server."
                )
            warnings.append(
                f"{len(scan_indexes)} halaman berupa gambar tidak dibaca (Tesseract OCR belum terpasang)."
            )
            for i in scan_indexes:
                pages[i].note = "[Halaman ini berupa gambar dan tidak dapat dibaca]"
        else:
            if len(scan_indexes) > MAX_OCR_PAGES:
                raise PdfProcessingError(
                    f"PDF scan memiliki {len(scan_indexes)} halaman. "
                    f"Maksimal {MAX_OCR_PAGES} halaman untuk OCR.\n\n"
                    "Gunakan Split PDF untuk memotongnya terlebih dahulu."
                )
            langs, missing = resolve_languages(languages, list_languages(tesseract))
            if missing:
                warnings.append(f"Bahasa OCR {', '.join(missing)} belum terpasang; memakai {langs}.")
            ocr_pages(input_path, scan_indexes, pages, work_dir, tesseract, langs)
            ocr_done = len(scan_indexes)

    if not any(p.blocks for p in pages):
        raise PdfProcessingError(
            "Tidak ada teks yang dapat dikenali di PDF ini.\n\n"
            "Pastikan dokumen terbaca jelas dan tidak kosong."
        )
    build_docx(pages, Path(output_path))
    return ConversionReport(kind=kind, page_count=len(pages), ocr_pages=ocr_done, warnings=warnings)
