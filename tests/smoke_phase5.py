"""Smoke test Phase 5 - TIDAK butuh Telegram/internet.

Bagian A: logika murni + program eksternal yang tersedia (LibreOffice / Tesseract
          dilewati otomatis jika belum terpasang).
Bagian B: uji nyata dengan PyMuPDF (dilewati jika belum terpasang).
Jalankan dari folder project:   python tests/smoke_phase5.py
"""
from __future__ import annotations

import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402
from pypdf import PdfReader  # noqa: E402

import services.pdf_to_word as p2w  # noqa: E402
import services.word_to_pdf as w2p  # noqa: E402
from services.pdf_common import MissingDependencyError, PdfProcessingError  # noqa: E402
from services.process_utils import ProcessTimeout, run_command  # noqa: E402

results: list[tuple[bool, str]] = []


def check(name: str, condition: bool) -> None:
    results.append((condition, name))
    print(("  PASS  " if condition else "  FAIL  ") + name)


def raises(exc_type, fn) -> bool:
    try:
        fn()
    except exc_type:
        return True
    return False


def docx_text(path: Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs]


def text_image(path: Path, lines: list[str]) -> bool:
    """Gambar teks besar untuk uji OCR. Return False jika font bawaan tidak bisa diatur ukurannya."""
    try:
        font = ImageFont.load_default(size=56)
    except TypeError:
        return False
    img = Image.new("RGB", (1400, 140 * len(lines) + 80), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((60, 40 + i * 140), line, fill="black", font=font)
    img.save(path)
    return True


def part_a(tmp: Path) -> None:
    print("\n[A1] process_utils")
    code, out, _ = run_command([sys.executable, "-c", "print('halo')"], timeout=20)
    check("perintah normal: kode 0 & stdout terbaca", code == 0 and out.strip() == b"halo")
    started = time.monotonic()
    timed_out = raises(ProcessTimeout, lambda: run_command([sys.executable, "-c", "import time; time.sleep(60)"], timeout=1))
    check("proses macet dihentikan paksa saat timeout (<8 detik)", timed_out and time.monotonic() - started < 8)

    print("\n[A2] bahasa & teks OCR")
    check("bahasa tersedia dipakai", p2w.resolve_languages("ind+eng", ("eng", "ind")) == ("ind+eng", []))
    check("bahasa hilang -> fallback + dilaporkan", p2w.resolve_languages("ind+eng", ("eng",)) == ("eng", ["ind"]))
    check("tidak ada yang cocok -> eng", p2w.resolve_languages("ind", ("eng", "fra"))[0] == "eng")
    check("tanpa data bahasa -> error", raises(PdfProcessingError, lambda: p2w.resolve_languages("ind", ())))
    check("paragraf OCR dipisah baris kosong, baris tunggal disambung",
          p2w.ocr_text_to_paragraphs("Baris satu\nlanjutan\n\n\nParagraf dua\x0c\n") == ["Baris satu lanjutan", "Paragraf dua"])
    check("kode bahasa berbahaya ditolak", raises(ValueError, lambda: p2w.run_tesseract(tmp / "x.png", "tesseract", "eng; rm -rf /")))

    print("\n[A3] ekstraksi blok & klasifikasi")
    raw = {"type": 0, "lines": [
        {"spans": [{"text": "Judul ", "size": 20.0, "flags": 16, "font": "Arial-Bold"},
                   {"text": "Besar", "size": 20.0, "flags": 16, "font": "Arial-Bold"}]},
        {"spans": [{"text": "baris kedua \x00dengan\x0b kontrol", "size": 11.0, "flags": 2, "font": "Arial"}]},
    ]}
    block = p2w.block_from_dict(raw)
    check("span berformat sama digabung, format terbaca", block.spans[0].text.strip() == "Judul Besar" and block.spans[0].bold and block.spans[0].size == 20.0)
    check("baris disambung spasi, karakter kontrol dibuang", block.text == "Judul Besar baris kedua dengan kontrol" and block.spans[-1].italic)
    check("blok kosong -> None", p2w.block_from_dict({"type": 0, "lines": [{"spans": [{"text": "  ", "size": 11}]}]}) is None)
    sp = lambda t: p2w.Block([p2w.Span(t)])
    text_page = p2w.PageContent(blocks=[sp("Ini halaman teks yang cukup panjang.")], has_images=True)
    scan_page = p2w.PageContent(blocks=[], has_images=True)
    blank_page = p2w.PageContent(blocks=[], has_images=False)
    check("halaman teks bukan scan", not text_page.needs_ocr)
    check("halaman tanpa teks + ada gambar = scan", scan_page.needs_ocr)
    check("halaman kosong tanpa gambar bukan scan", not blank_page.needs_ocr)
    check("klasifikasi text / scanned / mixed",
          p2w.classify([text_page, blank_page]) == "text" and p2w.classify([scan_page, scan_page]) == "scanned"
          and p2w.classify([text_page, scan_page]) == "mixed")

    print("\n[A4] build_docx")
    pages = [
        p2w.PageContent(width=595, height=842, blocks=[
            p2w.Block([p2w.Span("Judul Dokumen", 22, True)]),
            p2w.Block([p2w.Span("Isi paragraf pertama dengan ", 11), p2w.Span("tebal", 11, True), p2w.Span(" dan ", 11), p2w.Span("miring", 11, False, True)]),
        ]),
        p2w.PageContent(blocks=[p2w.Block([p2w.Span("Halaman dua.", 11)])]),
        p2w.PageContent(blocks=[], note="[Halaman ini berupa gambar dan tidak dapat dibaca]"),
    ]
    out = tmp / "built.docx"; p2w.build_docx(pages, out)
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    check("teks tiap halaman ada", "Judul Dokumen" in texts and "Isi paragraf pertama dengan tebal dan miring" in texts and "Halaman dua." in texts)
    para = next(p for p in doc.paragraphs if p.text.startswith("Isi paragraf"))
    check("format tebal & miring tersimpan per run", [r.bold for r in para.runs] == [False, True, False, False] and para.runs[3].italic is True)
    check("judul besar otomatis tebal", next(p for p in doc.paragraphs if p.text == "Judul Dokumen").runs[0].bold is True)
    check("page break antar halaman (2 pemisah)", out.read_bytes().count(b'w:br w:type="page"') == 2 or sum('w:br' in p._p.xml and 'type="page"' in p._p.xml for p in doc.paragraphs) == 2)
    check("catatan halaman gambar ikut ditulis", any("berupa gambar" in t for t in texts))
    check("ukuran halaman mengikuti PDF (A4)", round(doc.sections[0].page_width.pt) == 595)

    print("\n[A5] alur pdf_to_word (ekstraksi & OCR diganti tiruan)")
    real_extract, real_ocr, real_find = p2w.extract_pages, p2w.ocr_pages, p2w.find_tesseract
    try:
        p2w.extract_pages = lambda path: [text_page, blank_page]
        rep = p2w.pdf_to_word(tmp / "a.pdf", tmp / "t.docx", tmp)
        check("PDF teks: jenis 'text', tanpa OCR", rep.kind == "text" and rep.ocr_pages == 0 and (tmp / "t.docx").exists())

        mk_scan = lambda: p2w.PageContent(blocks=[], has_images=True)
        p2w.extract_pages = lambda path: [mk_scan(), mk_scan()]
        p2w.find_tesseract = lambda *_: None
        check("PDF scan tanpa Tesseract -> pesan jelas", raises(MissingDependencyError, lambda: p2w.pdf_to_word(tmp / "a.pdf", tmp / "s.docx", tmp)))

        p2w.extract_pages = lambda path: [text_page, mk_scan()]
        rep = p2w.pdf_to_word(tmp / "a.pdf", tmp / "m.docx", tmp)
        check("PDF campuran tanpa Tesseract -> tetap jadi, ada peringatan & catatan",
              rep.kind == "mixed" and rep.ocr_pages == 0 and rep.warnings and any("tidak dapat dibaca" in t for t in docx_text(tmp / "m.docx")))

        def fake_ocr(path, idx, pages, wd, cmd, langs):
            for i in idx:
                pages[i].blocks = [p2w.Block([p2w.Span("Hasil OCR halaman scan")])]
        p2w.find_tesseract = lambda *_: "/fake/tesseract"
        p2w.list_languages = lambda cmd: ("eng",)
        p2w.ocr_pages = fake_ocr
        p2w.extract_pages = lambda path: [mk_scan(), text_page]
        rep = p2w.pdf_to_word(tmp / "a.pdf", tmp / "o.docx", tmp, languages="ind+eng")
        check("PDF campuran dengan OCR: halaman scan terisi", rep.ocr_pages == 1 and "Hasil OCR halaman scan" in docx_text(tmp / "o.docx"))
        check("bahasa OCR yang hilang dilaporkan", any("ind" in w for w in rep.warnings))

        p2w.extract_pages = lambda path: [mk_scan() for _ in range(p2w.MAX_OCR_PAGES + 1)]
        check("terlalu banyak halaman scan ditolak", raises(PdfProcessingError, lambda: p2w.pdf_to_word(tmp / "a.pdf", tmp / "x.docx", tmp)))

        p2w.ocr_pages = lambda path, idx, pages, wd, cmd, langs: None   # OCR tidak menemukan apa-apa
        p2w.extract_pages = lambda path: [mk_scan()]
        check("hasil kosong sama sekali -> error jelas", raises(PdfProcessingError, lambda: p2w.pdf_to_word(tmp / "a.pdf", tmp / "e.docx", tmp)))
    finally:
        p2w.extract_pages, p2w.ocr_pages, p2w.find_tesseract = real_extract, real_ocr, real_find

    print("\n[A6] Tesseract nyata")
    tess = p2w.find_tesseract()
    if tess is None:
        print("  SKIP  Tesseract belum terpasang")
    else:
        img = tmp / "ocr.png"
        if text_image(img, ["Hello World", "Document Bot"]):
            text = p2w.run_tesseract(img, tess, p2w.resolve_languages("eng", p2w.list_languages(tess))[0])
            check(f"OCR membaca teks gambar: {text.split()[:3]}", "hello" in text.lower() and "world" in text.lower())
        else:
            print("  SKIP  Pillow lama: ukuran font bawaan tidak bisa diatur")

    print("\n[A7] LibreOffice nyata")
    office = w2p.find_libreoffice()
    check("find_libreoffice mengabaikan path palsu tanpa crash", w2p.find_libreoffice("/tidak/ada/soffice") == office)
    real_find_lo = w2p.find_libreoffice
    w2p.find_libreoffice = lambda *_: None
    try:
        check("LibreOffice tidak ada -> pesan jelas", raises(MissingDependencyError, lambda: w2p.word_to_pdf(tmp / "a.docx", tmp / "o.pdf")))
    finally:
        w2p.find_libreoffice = real_find_lo
    if office is None:
        print("  SKIP  LibreOffice belum terpasang")
        return
    src = tmp / "surat.docx"
    d = Document(); d.add_heading("Surat Resmi", 0); d.add_paragraph("Dengan hormat, ini adalah isi surat percobaan.")
    d.add_page_break(); d.add_paragraph("Halaman kedua."); d.save(str(src))
    out = tmp / "surat.pdf"
    pages_n = w2p.word_to_pdf(src, out, timeout=120)
    reader = PdfReader(str(out)); body = " ".join(p.extract_text() for p in reader.pages)
    check("DOCX -> PDF: 2 halaman, teks terbaca", pages_n == 2 and "Surat Resmi" in body and "Halaman kedua" in body)
    check("folder kerja LibreOffice dibersihkan", not any(p.name.startswith("lo_") for p in tmp.iterdir()))

    legacy_dir = tmp / "legacy"; legacy_dir.mkdir()
    from services.process_utils import run_command as rc
    rc([office, f"-env:UserInstallation={(tmp / 'prof').as_uri()}", "--headless", "--convert-to", "doc", "--outdir", str(legacy_dir), str(src)], timeout=120)
    legacy = legacy_dir / "surat.doc"
    if legacy.exists():
        pages_n = w2p.word_to_pdf(legacy, tmp / "legacy.pdf", timeout=120)
        check("DOC lama -> PDF juga berhasil", pages_n == 2)
    junk = tmp / "rusak.docx"; junk.write_bytes(b"PK\x03\x04" + b"\x00" * 200)
    try:
        w2p.word_to_pdf(junk, tmp / "rusak.pdf", timeout=60); handled = (tmp / "rusak.pdf").exists()
    except PdfProcessingError:
        handled = True
    check("dokumen rusak: tidak crash (error rapi atau tetap terkonversi)", handled)
    short = raises(PdfProcessingError, lambda: w2p.word_to_pdf(src, tmp / "t.pdf", timeout=0.05))
    check("timeout sangat pendek -> error rapi 'terlalu lama'", short)


def part_b(tmp: Path) -> None:
    from services.pdf_common import get_fitz
    try:
        fitz = get_fitz()
    except Exception:
        print("\n[B] SKIP - PyMuPDF belum terpasang (python -m pip install PyMuPDF)")
        return
    print(f"\n[B] Uji nyata dengan PyMuPDF {getattr(fitz, 'VersionBind', '?')}")

    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "Judul Laporan Percobaan", fontsize=22)
    page.insert_text((72, 140), "Ini adalah paragraf isi dokumen teks asli untuk pengujian konversi.", fontsize=11)
    text_pdf = tmp / "teks.pdf"; doc.save(str(text_pdf)); doc.close()
    out = tmp / "teks.docx"
    rep = p2w.pdf_to_word(text_pdf, out, tmp, languages="eng")
    body = " ".join(docx_text(out))
    check("PDF teks -> DOCX: jenis text, isi terbaca", rep.kind == "text" and "Judul Laporan" in body and "paragraf isi" in body)

    tess = p2w.find_tesseract()
    img = tmp / "scan.png"
    if tess and text_image(img, ["Hello World", "Scanned Page"]):
        doc = fitz.open(); page = doc.new_page(width=595, height=300)
        page.insert_image(page.rect, filename=str(img)); scan_pdf = tmp / "scan.pdf"; doc.save(str(scan_pdf)); doc.close()
        out = tmp / "scan.docx"
        rep = p2w.pdf_to_word(scan_pdf, out, tmp, languages="eng")
        check("PDF scan -> DOCX lewat OCR", rep.kind == "scanned" and rep.ocr_pages == 1 and "hello" in " ".join(docx_text(out)).lower())
    else:
        print("  SKIP  uji PDF scan (Tesseract / font bawaan tidak tersedia)")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="docbot_p5_"))
    part_a(tmp)
    part_b(tmp)
    failed = [name for ok, name in results if not ok]
    print(f"\nHasil: {len(results) - len(failed)}/{len(results)} lulus")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
