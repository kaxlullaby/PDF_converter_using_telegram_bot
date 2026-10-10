"""Pembuat file contoh untuk tes (PDF, gambar, DOCX, dan versi "jahat"-nya)."""
from __future__ import annotations

import io
import random
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pypdf import PdfWriter

MB = 1024 * 1024


def make_pdf(path: Path, widths: list[int] | tuple[int, ...] = (100,)) -> Path:
    """PDF dengan satu halaman per elemen `widths`; lebar halaman berbeda supaya urutannya bisa dibuktikan."""
    writer = PdfWriter()
    for width in widths:
        writer.add_blank_page(width=width, height=100)
    with Path(path).open("wb") as fh:
        writer.write(fh)
    return Path(path)


def page_widths(path_or_bytes) -> list[int]:
    from pypdf import PdfReader
    source = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else str(path_or_bytes)
    return [int(p.mediabox.width) for p in PdfReader(source).pages]


def make_png(path: Path, size=(50, 50), color="red", mode="RGB") -> Path:
    Image.new(mode, size, color).save(path)
    return Path(path)


def make_jpeg(path: Path, size=(400, 300), color="blue", orientation: int | None = None) -> Path:
    img = Image.new("RGB", size, color)
    if orientation is None:
        img.save(path, "JPEG")
    else:
        exif = img.getexif()
        exif[0x0112] = orientation
        img.save(path, "JPEG", exif=exif)
    return Path(path)


def noisy_image(w: int, h: int, seed: int = 1) -> Image.Image:
    """Gambar acak (sulit dikompres) supaya efek kompresi terlihat nyata."""
    rnd = random.Random(seed)
    base = Image.frombytes("RGB", (64, 64), bytes(rnd.getrandbits(8) for _ in range(64 * 64 * 3)))
    return base.resize((w, h), Image.BICUBIC).filter(ImageFilter.GaussianBlur(1))


def jpeg_bytes(img: Image.Image, quality: int = 95) -> bytes:
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=quality)
    return buf.getvalue()


def text_image(path: Path, lines: list[str]) -> bool:
    """Gambar teks besar untuk uji OCR. False jika Pillow lama tidak bisa mengatur ukuran font."""
    try:
        font = ImageFont.load_default(size=56)
    except TypeError:
        return False
    img = Image.new("RGB", (1400, 140 * len(lines) + 80), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((60, 40 + i * 140), line, fill="black", font=font)
    img.save(path)
    return True


CONTENT_TYPES = "<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'/>"


def rels_xml(*relationships: tuple[str, str, str | None]) -> str:
    """Relasi OOXML: (jenis, target, mode) -> XML. mode bisa 'External' atau None."""
    items = "".join(
        f'<Relationship Id="r{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}" '
        f'Target="{target}"' + (f' TargetMode="{mode}"' if mode else "") + "/>"
        for i, (kind, target, mode) in enumerate(relationships)
    )
    return f'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{items}</Relationships>'


def make_docx_zip(path: Path, extra: dict[str, bytes | str] | None = None, *, valid: bool = True) -> Path:
    """DOCX minimal buatan tangan (cukup untuk lolos validasi struktur). `extra` = berkas tambahan."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", CONTENT_TYPES)
        if valid:
            zf.writestr("word/document.xml", "<w:document xmlns:w='x'/>")
        for name, data in (extra or {}).items():
            zf.writestr(name, data)
    return Path(path)


def make_real_docx(path: Path, text: str = "Isi surat percobaan.") -> Path:
    from docx import Document
    doc = Document()
    doc.add_heading("Surat Resmi", 0)
    doc.add_paragraph(text)
    doc.save(str(path))
    return Path(path)
