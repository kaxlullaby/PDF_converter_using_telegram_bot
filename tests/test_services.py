"""Layanan pemrosesan: PDF (pypdf), gambar (Pillow), Word/OCR (logika + program eksternal bila ada)."""
from __future__ import annotations

import io
import unittest
import zipfile
from unittest.mock import patch

from docx import Document
from PIL import Image
from pypdf import PdfReader

import services.pdf_to_word as p2w
import services.word_to_pdf as w2p
from services.compress_pdf import PROFILES, recompress_image
from services.jpg_to_pdf import fit_rect, page_size_for, prepare_image
from services.merge_pdf import merge_pdfs
from services.pdf_common import MissingDependencyError, PdfProcessingError, get_fitz
from services.pdf_to_jpg import build_zip, compute_scale, delivery_mode
from services.rotate_pdf import rotate_pdf
from services.split_pdf import PageSelectionError, parse_page_ranges, split_pdf
from tests.fakes import TmpCase
from tests.samples import (jpeg_bytes, make_jpeg, make_pdf, make_png, make_real_docx, noisy_image,
                           page_widths, text_image)


class TestParsePageRanges(unittest.TestCase):
    def test_format_valid(self):
        cases = {"1-5": [1, 2, 3, 4, 5], "2,4,7": [2, 4, 7], "3-8": [3, 4, 5, 6, 7, 8], "1-3, 6, 9 - 10": [1, 2, 3, 6, 9, 10],
                 "5,1": [5, 1], "1,1-2": [1, 2], "3–4": [3, 4]}
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(parse_page_ranges(text, 12), expected)

    def test_format_salah(self):
        for text in ("", "0", "5-1", "13", "1-13", "abc", "1-", "-3", "1..3", "9" * 300, "1-999999999", "1;;x", "1/2"):
            with self.subTest(text=text[:15]):
                with self.assertRaises(PageSelectionError):
                    parse_page_ranges(text, 12)


class TestPdfServices(TmpCase):
    def setUp(self):
        super().setUp()
        self.a = make_pdf(self.tmp / "a.pdf", [101])
        self.b = make_pdf(self.tmp / "b.pdf", [201, 202])
        self.c = make_pdf(self.tmp / "c.pdf", [301, 302, 303])

    def test_merge_urutan_dan_jumlah(self):
        total = merge_pdfs([self.c, self.a, self.b], self.tmp / "m.pdf")
        self.assertEqual(total, 6)
        self.assertEqual(page_widths(self.tmp / "m.pdf"), [301, 302, 303, 101, 201, 202])

    def test_merge_minimal_dua(self):
        with self.assertRaises(PdfProcessingError):
            merge_pdfs([self.a], self.tmp / "x.pdf")

    def test_merge_batas_halaman_hasil(self):
        big = make_pdf(self.tmp / "big.pdf", [100] * 400)
        with self.assertRaises(PdfProcessingError):
            merge_pdfs([big] * 6, self.tmp / "x.pdf")  # 2400 > 2000

    def test_split(self):
        self.assertEqual(split_pdf(self.c, [3, 1], self.tmp / "s.pdf"), 2)
        self.assertEqual(page_widths(self.tmp / "s.pdf"), [303, 301])
        with self.assertRaises(PdfProcessingError):
            split_pdf(self.c, [4], self.tmp / "y.pdf")

    def test_rotate(self):
        rotate_pdf(self.c, 90, self.tmp / "r.pdf")
        self.assertEqual([p.rotation % 360 for p in PdfReader(str(self.tmp / "r.pdf")).pages], [90] * 3)
        rotate_pdf(self.tmp / "r.pdf", 270, self.tmp / "r2.pdf")  # rotasi menumpuk
        self.assertEqual({p.rotation % 360 for p in PdfReader(str(self.tmp / "r2.pdf")).pages}, {0})
        with self.assertRaises(ValueError):
            rotate_pdf(self.c, 45, self.tmp / "z.pdf")

    def test_pdf_terenkripsi_password_kosong_bisa_dipakai(self):
        from pypdf import PdfWriter
        writer = PdfWriter(); writer.add_blank_page(width=401, height=100); writer.encrypt(user_password="")
        with (self.tmp / "enc.pdf").open("wb") as fh:
            writer.write(fh)
        merge_pdfs([self.a, self.tmp / "enc.pdf"], self.tmp / "m.pdf")
        self.assertEqual(page_widths(self.tmp / "m.pdf"), [101, 401])


class TestImageServices(TmpCase):
    def test_recompress_level_dan_ukuran(self):
        src = jpeg_bytes(noisy_image(2400, 1800), 95)
        out = {lv: recompress_image(src, PROFILES[lv]) for lv in PROFILES}
        self.assertTrue(all(o is not None and len(o) < len(src) for o in out.values()))
        self.assertGreater(len(out["low"]), len(out["medium"]))
        self.assertGreater(len(out["medium"]), len(out["high"]))
        self.assertEqual(Image.open(io.BytesIO(out["low"])).size, (2400, 1800))
        self.assertLessEqual(max(Image.open(io.BytesIO(out["medium"])).size), 1800)
        self.assertLessEqual(max(Image.open(io.BytesIO(out["high"])).size), 1200)

    def test_recompress_dilewati_untuk_kasus_tidak_layak(self):
        self.assertIsNone(recompress_image(jpeg_bytes(Image.new("RGB", (40, 40), "red")), PROFILES["high"]))
        self.assertIsNone(recompress_image(b"x" * 50000, PROFILES["high"]))
        self.assertIsNone(recompress_image(jpeg_bytes(noisy_image(600, 400), 30), PROFILES["low"]))
        transparent = io.BytesIO(); Image.new("RGBA", (300, 300), (1, 2, 3, 4)).save(transparent, "PNG")
        self.assertIsNone(recompress_image(transparent.getvalue() + b"\0" * 30000, PROFILES["high"]))

    def test_recompress_draft_menjaga_kualitas_dan_bentuk(self):
        """Optimasi decode JPEG dalam ukuran kecil tidak boleh mengubah dimensi / rasio hasil akhir."""
        src = jpeg_bytes(noisy_image(4000, 3000), 92)
        for level, side in (("medium", 1800), ("high", 1200)):
            image = Image.open(io.BytesIO(recompress_image(src, PROFILES[level])))
            self.assertEqual(image.size, (side, side * 3 // 4))

    def test_jpg_to_pdf_persiapan_gambar(self):
        self.assertEqual(page_size_for(300, 400), (595.0, 842.0))
        self.assertEqual(page_size_for(400, 300), (842.0, 595.0))
        x0, y0, x1, y1 = fit_rect(1000, 500, 595, 842, 20)
        self.assertAlmostEqual(x0, 20); self.assertAlmostEqual(x1, 575); self.assertAlmostEqual((y0 + y1) / 2, 421)

        plain = make_jpeg(self.tmp / "p.jpg")
        data, w, h = prepare_image(plain)
        self.assertEqual((data, w, h), (plain.read_bytes(), 400, 300))  # tanpa kompres ulang

        data, w, h = prepare_image(make_png(self.tmp / "t.png", mode="RGBA", color=(255, 0, 0, 0)))
        self.assertGreater(min(Image.open(io.BytesIO(data)).convert("RGB").getpixel((25, 25))), 240)  # latar putih

        data, w, h = prepare_image(make_jpeg(self.tmp / "r.jpg", (400, 200), orientation=6))
        self.assertEqual((w, h), (200, 400))  # foto HP miring ditegakkan

        Image.new("RGB", (6000, 3000), "white").save(self.tmp / "huge.png")
        data, w, h = prepare_image(self.tmp / "huge.png")
        self.assertEqual(max(w, h), 3000)

    def test_pdf_to_jpg_logika(self):
        self.assertAlmostEqual(compute_scale(595, 842), 150 / 72)
        self.assertAlmostEqual(5000 * compute_scale(5000, 3000), 3000)
        self.assertGreater(compute_scale(0, 0), 0)
        self.assertEqual([delivery_mode(n) for n in (1, 2, 10, 11)], ["single", "multiple", "multiple", "zip"])
        files = []
        for i in range(3):
            p = self.tmp / f"p{i}.jpg"; Image.new("RGB", (20, 20), "red").save(p); files.append((p, f"d_{i + 1:02d}.jpg"))
        build_zip(files, self.tmp / "z.zip")
        self.assertEqual(zipfile.ZipFile(self.tmp / "z.zip").namelist(), ["d_01.jpg", "d_02.jpg", "d_03.jpg"])


class TestPdfToWordLogic(TmpCase):
    """Logika PDF -> Word dengan ekstraksi/OCR diganti tiruan (PyMuPDF diuji terpisah)."""

    def setUp(self):
        super().setUp()
        sp = lambda t: p2w.Block([p2w.Span(t)])
        self.text_page = p2w.PageContent(blocks=[sp("Ini halaman teks yang cukup panjang.")], has_images=True)
        self.scan = lambda: p2w.PageContent(blocks=[], has_images=True)
        self.blank = p2w.PageContent(blocks=[], has_images=False)

    def convert(self, pages, **kwargs):
        with patch.object(p2w, "extract_pages", lambda path: pages):
            return p2w.pdf_to_word(self.tmp / "a.pdf", self.tmp / "out.docx", self.tmp, **kwargs)

    def texts(self):
        return [p.text for p in Document(str(self.tmp / "out.docx")).paragraphs]

    def test_blok_dan_klasifikasi(self):
        raw = {"type": 0, "lines": [
            {"spans": [{"text": "Judul ", "size": 20.0, "flags": 16, "font": "Arial-Bold"}, {"text": "Besar", "size": 20.0, "flags": 16, "font": "Arial-Bold"}]},
            {"spans": [{"text": "baris kedua \x00dengan\x0b kontrol", "size": 11.0, "flags": 2, "font": "Arial"}]}]}
        block = p2w.block_from_dict(raw)
        self.assertEqual(block.text, "Judul Besar baris kedua dengan kontrol")
        self.assertTrue(block.spans[0].bold and block.spans[-1].italic)
        self.assertIsNone(p2w.block_from_dict({"type": 0, "lines": [{"spans": [{"text": "  ", "size": 11}]}]}))
        self.assertTrue(self.scan().needs_ocr)
        self.assertFalse(self.blank.needs_ocr or self.text_page.needs_ocr)
        self.assertEqual(p2w.classify([self.text_page, self.blank]), "text")
        self.assertEqual(p2w.classify([self.scan(), self.scan()]), "scanned")
        self.assertEqual(p2w.classify([self.text_page, self.scan()]), "mixed")

    def test_bahasa_dan_paragraf_ocr(self):
        self.assertEqual(p2w.resolve_languages("ind+eng", ("eng", "ind")), ("ind+eng", []))
        self.assertEqual(p2w.resolve_languages("ind+eng", ("eng",)), ("eng", ["ind"]))
        with self.assertRaises(PdfProcessingError):
            p2w.resolve_languages("ind", ())
        self.assertEqual(p2w.ocr_text_to_paragraphs("Baris satu\nlanjutan\n\n\nParagraf dua\x0c\n"), ["Baris satu lanjutan", "Paragraf dua"])
        with self.assertRaises(ValueError):
            p2w.run_tesseract(self.tmp / "x.png", "tesseract", "eng; rm -rf /")

    def test_build_docx_format(self):
        pages = [p2w.PageContent(blocks=[p2w.Block([p2w.Span("Judul Dokumen", 22, True)]),
                                         p2w.Block([p2w.Span("Isi ", 11), p2w.Span("tebal", 11, True), p2w.Span(" dan ", 11), p2w.Span("miring", 11, False, True)])]),
                 p2w.PageContent(blocks=[p2w.Block([p2w.Span("Halaman dua.", 11)])]),
                 p2w.PageContent(note="[Halaman ini berupa gambar dan tidak dapat dibaca]")]
        p2w.build_docx(pages, self.tmp / "out.docx")
        doc = Document(str(self.tmp / "out.docx"))
        para = next(p for p in doc.paragraphs if p.text.startswith("Isi"))
        self.assertEqual([r.bold for r in para.runs], [False, True, False, False])
        self.assertTrue(para.runs[3].italic)
        self.assertTrue(next(p for p in doc.paragraphs if p.text == "Judul Dokumen").runs[0].bold)
        self.assertEqual(sum('type="page"' in p._p.xml for p in doc.paragraphs), 2)
        self.assertEqual(round(doc.sections[0].page_width.pt), 595)

    def test_alur_tanpa_ocr(self):
        report = self.convert([self.text_page, self.blank])
        self.assertEqual((report.kind, report.ocr_pages), ("text", 0))

    def test_scan_tanpa_tesseract(self):
        with patch.object(p2w, "find_tesseract", lambda *_: None):
            with self.assertRaises(MissingDependencyError):
                self.convert([self.scan(), self.scan()])
            report = self.convert([self.text_page, self.scan()])
        self.assertEqual((report.kind, report.ocr_pages), ("mixed", 0))
        self.assertTrue(report.warnings and any("tidak dapat dibaca" in t for t in self.texts()))

    def test_scan_dengan_ocr_tiruan(self):
        def fake_ocr(path, idx, pages, wd, cmd, langs):
            for i in idx:
                pages[i].blocks = [p2w.Block([p2w.Span("Hasil OCR halaman scan")])]
        with patch.object(p2w, "find_tesseract", lambda *_: "/fake"), patch.object(p2w, "list_languages", lambda c: ("eng",)), \
                patch.object(p2w, "ocr_pages", fake_ocr):
            report = self.convert([self.scan(), self.text_page], languages="ind+eng")
            self.assertEqual(report.ocr_pages, 1)
            self.assertIn("Hasil OCR halaman scan", self.texts())
            self.assertTrue(any("ind" in w for w in report.warnings))
            with self.assertRaises(PdfProcessingError):  # terlalu banyak halaman scan
                self.convert([self.scan() for _ in range(p2w.MAX_OCR_PAGES + 1)])
        with patch.object(p2w, "find_tesseract", lambda *_: "/fake"), patch.object(p2w, "list_languages", lambda c: ("eng",)), \
                patch.object(p2w, "ocr_pages", lambda *a: None):
            with self.assertRaises(PdfProcessingError):  # OCR tidak menemukan teks apa pun
                self.convert([self.scan()])


@unittest.skipUnless(p2w.find_tesseract(), "Tesseract belum terpasang")
class TestTesseractReal(TmpCase):
    def test_ocr_membaca_teks_gambar(self):
        image = self.tmp / "ocr.png"
        if not text_image(image, ["Hello World", "Document Bot"]):
            self.skipTest("Pillow lama: ukuran font bawaan tidak bisa diatur")
        tess = p2w.find_tesseract()
        langs, _ = p2w.resolve_languages("eng", p2w.list_languages(tess))
        text = p2w.run_tesseract(image, tess, langs).lower()
        self.assertIn("hello", text); self.assertIn("world", text)


@unittest.skipUnless(w2p.find_libreoffice(), "LibreOffice belum terpasang")
class TestLibreOfficeReal(TmpCase):
    def test_docx_ke_pdf(self):
        src = make_real_docx(self.tmp / "surat.docx")
        pages = w2p.word_to_pdf(src, self.tmp / "surat.pdf", timeout=120)
        body = " ".join(p.extract_text() for p in PdfReader(str(self.tmp / "surat.pdf")).pages)
        self.assertEqual(pages, 1); self.assertIn("Surat Resmi", body); self.assertIn("percobaan", body)
        self.assertFalse(any(p.name.startswith("lo_") for p in self.tmp.iterdir()))  # folder kerja bersih

    def test_timeout_dan_dokumen_rusak(self):
        src = make_real_docx(self.tmp / "surat.docx")
        with self.assertRaises(PdfProcessingError):
            w2p.word_to_pdf(src, self.tmp / "t.pdf", timeout=0.05)
        junk = self.tmp / "rusak.docx"; junk.write_bytes(b"PK\x03\x04" + b"\x00" * 200)
        try:
            w2p.word_to_pdf(junk, self.tmp / "rusak.pdf", timeout=60)
        except PdfProcessingError:
            pass  # error rapi = diterima; tidak boleh crash/hang

    def test_libreoffice_hilang(self):
        with patch.object(w2p, "find_libreoffice", lambda *_: None):
            with self.assertRaises(MissingDependencyError):
                w2p.word_to_pdf(self.tmp / "a.docx", self.tmp / "o.pdf")
        self.assertEqual(w2p.find_libreoffice("/tidak/ada/soffice"), w2p.find_libreoffice())


def _has_fitz() -> bool:
    try:
        get_fitz(); return True
    except PdfProcessingError:
        return False


@unittest.skipUnless(_has_fitz(), "PyMuPDF belum terpasang (python -m pip install PyMuPDF)")
class TestPyMuPDFIntegration(TmpCase):
    """Uji NYATA dengan PyMuPDF: compress, render PDF -> JPG, JPG -> PDF, PDF -> Word."""

    def setUp(self):
        super().setUp()
        self.fitz = get_fitz()
        doc = self.fitz.open()
        for i in range(3):
            page = doc.new_page(width=595, height=842)
            page.insert_text((72, 72), f"Halaman {i + 1} Judul Laporan", fontsize=20)
            page.insert_image(self.fitz.Rect(50, 100, 545, 500), stream=jpeg_bytes(noisy_image(1600, 1300, seed=i), 95))
        self.sample = self.tmp / "sample.pdf"
        doc.save(str(self.sample)); doc.close()

    def test_compress(self):
        from services.compress_pdf import compress_pdf
        sizes = {}
        for level in PROFILES:
            original, compressed = compress_pdf(self.sample, level, self.tmp / f"c_{level}.pdf")
            with self.fitz.open(str(self.tmp / f"c_{level}.pdf")) as doc:
                self.assertEqual(doc.page_count, 3)
            self.assertLessEqual(compressed, original)
            sizes[level] = compressed
        self.assertLessEqual(sizes["high"], sizes["medium"]); self.assertLessEqual(sizes["medium"], sizes["low"])
        self.assertLess(sizes["high"], self.sample.stat().st_size * 0.8)

    def test_pdf_ke_jpg(self):
        from services.pdf_to_jpg import render_pages
        out = self.tmp / "jpgs"; out.mkdir()
        pages = render_pages(self.sample, out)
        self.assertEqual(len(pages), 3)
        self.assertTrue(all(Image.open(p).format == "JPEG" for p in pages))
        self.assertLessEqual(abs(Image.open(pages[0]).size[0] - 1240), 2)

    def test_jpg_ke_pdf(self):
        from services.jpg_to_pdf import images_to_pdf
        imgs = [make_jpeg(self.tmp / "i1.jpg", (800, 1200)), make_png(self.tmp / "i2.png", (1200, 600), (0, 0, 255, 128), "RGBA"),
                make_jpeg(self.tmp / "i3.jpg", (500, 500))]
        self.assertEqual(images_to_pdf(imgs, self.tmp / "o.pdf"), 3)
        with self.fitz.open(str(self.tmp / "o.pdf")) as doc:
            self.assertEqual([(round(p.rect.width), round(p.rect.height)) for p in doc], [(595, 842), (842, 595), (595, 842)])

    def test_pdf_ke_word_teks(self):
        report = p2w.pdf_to_word(self.sample, self.tmp / "o.docx", self.tmp, languages="eng")
        self.assertEqual(report.kind, "text")
        self.assertIn("Judul Laporan", " ".join(p.text for p in Document(str(self.tmp / "o.docx")).paragraphs))

    @unittest.skipUnless(p2w.find_tesseract(), "Tesseract belum terpasang")
    def test_pdf_scan_ke_word_dengan_ocr(self):
        image = self.tmp / "scan.png"
        if not text_image(image, ["Hello World", "Scanned Page"]):
            self.skipTest("font bawaan tidak bisa diatur")
        doc = self.fitz.open(); page = doc.new_page(width=595, height=300)
        page.insert_image(page.rect, filename=str(image)); doc.save(str(self.tmp / "scan.pdf")); doc.close()
        report = p2w.pdf_to_word(self.tmp / "scan.pdf", self.tmp / "s.docx", self.tmp, languages="eng")
        self.assertEqual((report.kind, report.ocr_pages), ("scanned", 1))
        self.assertIn("hello", " ".join(p.text for p in Document(str(self.tmp / "s.docx")).paragraphs).lower())


if __name__ == "__main__":
    unittest.main()
